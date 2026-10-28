"""Importação das propostas da biblioteca do SharePoint (só leitura).

Estrutura esperada (site de cada empresa, ver ``Empresa.sharepoint_*``; uma ou mais pastas de
propostas, cada uma com uma faixa de anos opcional, ex.: ``Propostas_antigas: 2016-2020``)::

    <pasta das propostas>/<ano>/<cliente>/Proposta 021 26.docx
                                          Proposta 021 26 Rev 01.docx / .pdf
                                          Pedido Compra 4100089244 - prop_021 26.pdf
    <pasta de controle>/Propostas Treinamento e consultoria2026.xlsx

Regras (detalhes em ``documents/sharepoint.md``):

* cada número (``Proposta NNN AA``) é uma proposta, com o número e o ano originais; a maior
  revisão (``Rev NN``) é o documento principal e os dados saem do .docx dela;
* a planilha de controle do ano completa serviço, cliente e data;
* pedido de compra na pasta do cliente que cita o número da proposta, ou numa pasta com uma
  só proposta, marca a proposta como ganha; sem pedido e com a validade vencida, expirada;
* todos os arquivos dessas pastas (e das pastas de contratos) entram no acervo, classificados
  pelo nome (proposta, pedido, contrato, NDA, cadastro, fiscal, dados, outro);
* proposta criada no sistema com o mesmo número não é tocada (vira conflito no relatório);
  na importada, só os dados do documento são atualizados quando ele muda; etapa decidida no
  sistema (negociação, ganha, perdida, cancelada) não é desfeita.
"""

import io
import re
from contextlib import nullcontext
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal

from django.contrib.auth.models import User
from django.db import transaction
from django.utils import timezone

from .constantes import ENVIADA, EXPIRADA, GANHA, PROBABILIDADE_PADRAO, ROTULOS_ETAPA, ZERO
from .leitor_docx import DadosDocumento, ler
from .models import (
    ArquivoProposta,
    Cliente,
    Contato,
    Empresa,
    EventoProposta,
    ArquivoSharePoint,
    ContratoCliente,
    ExecucaoSincronizacao,
    ItemProposta,
    Proposta,
)
from .servicos import recalcular, registrar
from .sharepoint import ItemRemoto, NaoEncontrado, SharePointClient, SharePointErro

# "Proposta 021 26", "Proposta 186-21", "Proposta Comercial 170 24 rev01", "Proposta Técnica e Comercial 147 26".
ARQUIVO_PROPOSTA = re.compile(
    r"^proposta(?:[\s_]+(?P<tipo>t[ée]cnica(?:[\s_]+e[\s_]+comercial)?|comercial))?"
    r"[\s_-]*(?P<numero>\d{1,4})[\s_/.-]+(?P<aa>\d{2})(?!\d)(?:[\s_-]*rev(?:is[aã]o)?[\s_.-]*(?P<rev>\d+))?",
    re.I,
)
ARQUIVO_PEDIDO = re.compile(r"pedido|^po[\s_-]|purchase|ordem de compra|folha de servi", re.I)
EXTENSOES_DOCUMENTO = (".docx", ".pptx", ".pdf", ".doc")
EXTENSOES_LIDAS = (".docx", ".pptx")
PROFUNDIDADE_NO_ANO = 5  # ano/cliente/sistema/versão/arquivo
FAIXA_DE_ANOS = re.compile(r"^(?P<pasta>.+?)\s*:\s*(?P<inicio>\d{4})?\s*-\s*(?P<fim>\d{4})?\s*$")
# Ordem importa: o primeiro padrão que casa define a categoria do arquivo no acervo.
CATEGORIAS_ACERVO = [
    (ArquivoSharePoint.NDA, re.compile(r"\bnda\b|confidencial|sigilo|non.?disclosure", re.I)),
    (ArquivoSharePoint.CONTRATO, re.compile(r"contrat|aditivo|minuta|clicksign|docusign|assinad", re.I)),
    (ArquivoSharePoint.PEDIDO, ARQUIVO_PEDIDO),
    (ArquivoSharePoint.CADASTRO, re.compile(r"cadastr|fornecedor|vendor|ariba|certid|\bcnd\b|cnpj|ficha|alvar|balan[cç]o", re.I)),
    (ArquivoSharePoint.FISCAL, re.compile(r"\bnf|nota fiscal|nfe|danfe|fatura|boleto|recibo|comprovante", re.I)),
]


def pastas_propostas(valor: str) -> list[tuple[str, int | None, int | None]]:
    """"Propostas_antigas: 2016-2020" -> ("Propostas_antigas", 2016, 2020); sem faixa, todos os anos."""
    pastas = []
    for linha in (valor or "").splitlines():
        linha = linha.strip()
        if not linha:
            continue
        m = FAIXA_DE_ANOS.match(linha)
        if m:
            inicio, fim = m["inicio"], m["fim"]
            pastas.append((m["pasta"].strip().strip("/"), int(inicio) if inicio else None, int(fim) if fim else None))
        else:
            pastas.append((linha.strip("/"), None, None))
    return pastas


def classificar(nome: str) -> str:
    if ARQUIVO_PROPOSTA.match(nome):
        return ArquivoSharePoint.PROPOSTA
    categoria = next((c for c, padrao in CATEGORIAS_ACERVO if padrao.search(nome)), None)
    if categoria:
        return categoria
    if nome.lower().endswith((".xls", ".xlsx", ".xlsm", ".csv")):
        return ArquivoSharePoint.DADOS
    return ArquivoSharePoint.OUTRO


def juntar(destino: dict, grupos: dict) -> None:
    """Junta os grupos de outra pasta: a mesma proposta em duas pastas vira uma só, sem arquivos repetidos."""
    for chave, grupo in grupos.items():
        atual = destino.setdefault(chave, grupo)
        if atual is grupo:
            continue
        vistos = {(rev, item.nome.lower()) for rev, item in atual.documentos}
        atual.documentos += [(rev, item) for rev, item in grupo.documentos if (rev, item.nome.lower()) not in vistos]
        pedidos = {item.nome.lower() for item in atual.pedidos}
        atual.pedidos += [item for item in grupo.pedidos if item.nome.lower() not in pedidos]
        atual.pasta_url = atual.pasta_url or grupo.pasta_url
NUMERO_PLANILHA = re.compile(r"(\d{1,4})\s*[/ ]\s*(\d{2})")


def _normal(texto: str) -> str:
    base = unicodedata.normalize("NFKD", (texto or "").replace("_", " ").replace("\xa0", " "))
    return " ".join(base.encode("ascii", "ignore").decode().lower().split())


def cita_numero(nome: str, numero: int, ano: int) -> bool:
    """O nome do arquivo cita a proposta (ex.: "prop_002 26", "039 24", "021/26")."""
    return bool(re.search(rf"(?<!\d)0*{numero}\s*[/ _.-]\s*{ano % 100:02d}(?!\d)", nome))


@dataclass
class Grupo:
    """Arquivos de uma proposta no SharePoint."""

    ano: int
    numero: int
    cliente: str
    pasta_url: str = ""
    documentos: list[tuple[int, ItemRemoto]] = field(default_factory=list)
    pedidos: list[ItemRemoto] = field(default_factory=list)

    @property
    def revisao(self) -> int:
        return max(rev for rev, _ in self.documentos)

    @property
    def principal(self) -> ItemRemoto:
        """Documento da maior revisão de onde os dados são lidos: o comercial (a proposta só técnica não tem preço),
        em .docx ou .pptx, o mais recente."""
        da_revisao = [item for rev, item in self.documentos if rev == self.revisao]

        def ordem(item: ItemRemoto):
            so_tecnica = bool(re.match(r"proposta[\s_]+t[ée]cnica(?![\s_]+e[\s_]+comercial)", item.nome, re.I))
            return (not item.nome.lower().endswith(EXTENSOES_LIDAS), so_tecnica, -(item.modificado.timestamp() if item.modificado else 0))

        return min(da_revisao, key=ordem)


@dataclass(frozen=True)
class LinhaControle:
    servico: str
    cliente: str
    data: date | None
    solicitante: str
    enviado_por: str = ""


# Colunas das planilhas de controle (nome normalizado → campo); os nomes mudaram ao longo dos anos.
COLUNAS_CONTROLE = {
    "servico": "servico", "sistema": "servico",
    "empresa": "cliente", "cliente": "cliente",
    "data": "data",
    "responsavel": "solicitante", "responsavel da equipe do projeto": "solicitante",
    "enviado por": "enviado_por", "enviado por:": "enviado_por",
}
LINHAS_VAZIAS_PARA_PARAR = 300  # há planilhas formatadas até a linha 1.048.576


@dataclass
class Relatorio:
    criadas: int = 0
    atualizadas: int = 0
    sem_mudanca: int = 0
    ganhas: int = 0
    expiradas: int = 0
    arquivos: int = 0
    contratos: int = 0
    conflitos: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    def resumo(self) -> str:
        return (
            f"{self.criadas} criada(s), {self.atualizadas} atualizada(s), {self.sem_mudanca} sem mudança; "
            f"{self.ganhas} marcada(s) como ganha(s) por pedido de compra, {self.expiradas} expirada(s); "
            f"{self.arquivos} arquivo(s) no acervo; "
            + (f"{self.contratos} contrato(s) na carteira; " if self.contratos else "")
            + f"{len(self.conflitos)} conflito(s), {len(self.avisos)} aviso(s)."
        )


def agrupar(itens: list[ItemRemoto], raiz: str, desde: int, ate: int | None = None) -> dict[tuple[int, int], Grupo]:
    """Propostas (por ano e número) a partir da listagem de uma pasta de propostas."""
    raiz = raiz.strip("/")
    urls_pastas = {i.caminho: i.url for i in itens if i.pasta}
    grupos: dict[tuple[int, int], Grupo] = {}
    pedidos_por_pasta: dict[tuple[int, str], list[ItemRemoto]] = defaultdict(list)
    for item in itens:
        partes = item.caminho[len(raiz):].strip("/").split("/") if raiz else item.caminho.split("/")
        if item.pasta or len(partes) < 2 or not partes[0].isdigit() or int(partes[0]) < desde:
            continue
        if ate is not None and int(partes[0]) > ate:
            continue
        ano = int(partes[0])
        cliente = partes[1] if len(partes) >= 3 else ""
        m = ARQUIVO_PROPOSTA.match(item.nome)
        if m and int(m["aa"]) == ano % 100 and item.nome.lower().endswith(EXTENSOES_DOCUMENTO):
            chave = (ano, int(m["numero"]))
            pasta = "/".join(filter(None, [raiz, partes[0], cliente]))
            grupo = grupos.setdefault(chave, Grupo(ano, chave[1], cliente, urls_pastas.get(pasta, "")))
            grupo.documentos.append((int(m["rev"] or 0), item))
        elif ARQUIVO_PEDIDO.search(item.nome) and cliente:
            pedidos_por_pasta[(ano, cliente)].append(item)

    for (ano, cliente), pedidos in pedidos_por_pasta.items():
        da_pasta = [g for g in grupos.values() if g.ano == ano and g.cliente == cliente]
        for pedido in pedidos:
            citados = [g for g in da_pasta if cita_numero(pedido.nome, g.numero, ano)]
            for grupo in citados or (da_pasta if len(da_pasta) == 1 else []):
                grupo.pedidos.append(pedido)
    return grupos


def ler_controle(conteudo: bytes, ano: int) -> dict[int, LinhaControle]:
    """Linhas da planilha de controle do ano, pelo número da proposta."""
    from openpyxl import load_workbook

    aba = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
    linhas = aba.iter_rows(values_only=True)
    colunas: dict[str, int] = {}
    for cabecalho in linhas:
        nomes = [_normal(str(c or "")) for c in cabecalho]
        if "n" in nomes or "no" in nomes or "numero" in nomes:
            colunas = {nome: i for i, nome in enumerate(nomes) if nome}
            break
    numero_col = next((colunas[n] for n in ("n", "no", "numero") if n in colunas), None)
    if numero_col is None:
        return {}

    indices = {campo: i for nome, i in colunas.items() if (campo := COLUNAS_CONTROLE.get(nome.rstrip(":")))}

    def celula(linha, campo):
        indice = indices.get(campo)
        return linha[indice] if indice is not None and indice < len(linha) else None

    resultado = {}
    vazias = 0
    for linha in linhas:
        if not any(v not in (None, "") for v in linha):
            vazias += 1
            if vazias >= LINHAS_VAZIAS_PARA_PARAR:
                break
            continue
        vazias = 0
        m = NUMERO_PLANILHA.search(str(linha[numero_col] or "")) if numero_col < len(linha) else None
        if not m or int(m.group(2)) != ano % 100:
            continue
        data = celula(linha, "data")
        if isinstance(data, datetime):
            data = data.date()
        elif isinstance(data, str):
            data = next((d for f in ("%d/%m/%Y", "%d/%m/%y") if (d := _data(data, f))), None)
        resultado[int(m.group(1))] = LinhaControle(
            servico=str(celula(linha, "servico") or "").strip(),
            cliente=str(celula(linha, "cliente") or "").strip(),
            data=data if isinstance(data, date) else None,
            solicitante=str(celula(linha, "solicitante") or "").strip(),
            enviado_por=str(celula(linha, "enviado_por") or "").strip(),
        )
    return resultado


def _data(texto: str, formato: str) -> date | None:
    try:
        return datetime.strptime(texto.strip(), formato).date()
    except ValueError:
        return None


class Sincronizacao:
    def __init__(self, empresa: Empresa, cliente: SharePointClient, hoje: date | None = None) -> None:
        self.empresa = empresa
        self.sp = cliente
        self.hoje = hoje or timezone.localdate()
        self.relatorio = Relatorio()
        clientes = list(Cliente.objects.filter(empresa=empresa))
        self._clientes = {_normal(str(c)): c for c in clientes}
        self._clientes.update({_normal(c.nome): c for c in clientes})
        # Pastas já ligadas valem mais que o nome, que pode ter sido mudado no sistema.
        self._clientes.update({_normal(n): c for c in clientes for n in c.nomes_sharepoint.splitlines() if n.strip()})
        self._usuarios = list(User.objects.filter(is_active=True))

    # -------------------------------------------------------------- leitura do SharePoint

    def listar(self) -> tuple[dict[tuple[int, int], Grupo], dict[tuple[int, int], LinhaControle]]:
        """Propostas de todas as pastas configuradas; guarda em ``self.acervo`` os arquivos vistos."""
        grupos: dict[tuple[int, int], Grupo] = {}
        self.acervo: list[tuple[ItemRemoto, str, int | None, str]] = []
        todos_os_anos: set[int] = set()
        for raiz, inicio, fim in pastas_propostas(self.empresa.sharepoint_pasta):
            desde = max(inicio or 0, self.empresa.sharepoint_desde)
            try:
                anos = [
                    i for i in self.sp.listar(raiz)
                    if i.pasta and i.nome.isdigit() and desde <= int(i.nome) <= (fim or 9999)
                ]
            except NaoEncontrado as erro:
                self.relatorio.avisos.append(f"Pasta “{raiz}” não lida: {erro}")
                continue
            itens = list(anos)
            for ano in anos:
                itens.extend(self.sp.listar(ano.caminho, profundidade=PROFUNDIDADE_NO_ANO))
            juntar(grupos, agrupar(itens, raiz, desde, fim))
            todos_os_anos.update(int(a.nome) for a in anos)
            for item in itens:
                partes = item.caminho[len(raiz):].strip("/").split("/")
                if not item.pasta and partes[0].isdigit():
                    self.acervo.append((item, raiz, int(partes[0]), partes[1] if len(partes) >= 3 else ""))
        for pasta in filter(None, (linha.strip().strip("/") for linha in self.empresa.sharepoint_contratos.splitlines())):
            try:
                itens = self.sp.listar(pasta, profundidade=3)
            except NaoEncontrado as erro:
                self.relatorio.avisos.append(f"Pasta de contratos “{pasta}” não lida: {erro}")
                continue
            for item in itens:
                if not item.pasta:
                    ano = re.search(r"(20\d{2})", item.caminho)
                    self.acervo.append((item, pasta, int(ano.group(1)) if ano else None, ""))
        return grupos, self._controle(todos_os_anos)

    def _controle(self, anos: set[int]) -> dict[tuple[int, int], LinhaControle]:
        pasta = self.empresa.sharepoint_controle.strip("/")
        if not pasta:
            return {}
        try:
            planilhas = [i for i in self.sp.listar(pasta) if i.nome.lower().endswith(".xlsx")]
        except SharePointErro as erro:
            self.relatorio.avisos.append(f"Planilhas de controle não lidas: {erro}")
            return {}
        controle = {}
        for ano in sorted(anos):
            # A planilha do ano termina com o ano (cópias como "2024-rubens-estat" ficam de fora).
            planilha = next((p for p in planilhas if re.search(rf"{ano}\.xlsx$", p.nome)), None)
            if planilha is None:
                continue
            try:
                linhas = ler_controle(self.sp.baixar(planilha), ano)
            except Exception as erro:  # planilha corrompida não impede a importação dos documentos
                self.relatorio.avisos.append(f"{planilha.nome}: não foi possível ler ({erro}).")
                continue
            controle.update({(ano, numero): linha for numero, linha in linhas.items()})
        return controle

    # -------------------------------------------------------------- gravação

    def _cliente(self, pasta: str) -> Cliente:
        """Cliente da pasta; a pasta fica registrada nele para as próximas sincronizações."""
        nome = " ".join(pasta.replace("_", " ").split()) or "Cliente não identificado"
        chave = _normal(nome)
        cliente = self._clientes.get(chave) or Cliente.objects.create(empresa=self.empresa, nome=nome)
        self._clientes[chave] = cliente
        if pasta and chave not in {_normal(n) for n in cliente.nomes_sharepoint.splitlines()}:
            cliente.nomes_sharepoint = "\n".join(filter(None, [cliente.nomes_sharepoint.strip(), pasta.strip()]))
            cliente.save(update_fields=["nomes_sharepoint"])
        return cliente

    def _contato(self, cliente: Cliente, nome: str) -> Contato | None:
        nome = nome.strip()
        if not nome:
            return None
        existente = next((c for c in cliente.contatos.all() if _normal(c.nome) == _normal(nome)), None)
        return existente or Contato.objects.create(cliente=cliente, nome=nome[:150])

    def _responsavel(self, assinante: str) -> User | None:
        """Usuário cujo nome (ou o início do e-mail) aparece na assinatura do documento."""
        palavras = set(_normal(assinante).split())
        if not palavras:
            return None
        for usuario in self._usuarios:
            nome = set(_normal(usuario.get_full_name()).split())
            login = _normal(usuario.email.split("@")[0]) if usuario.email else ""
            if (nome and nome <= palavras) or (login and login in palavras):
                return usuario
        return None

    def _dados(self, grupo: Grupo) -> DadosDocumento:
        principal = grupo.principal
        if not principal.nome.lower().endswith(EXTENSOES_LIDAS):
            self.relatorio.avisos.append(f"{self._rotulo(grupo)}: sem .docx/.pptx da última revisão; dados só da planilha.")
            return DadosDocumento()
        try:
            return ler(self.sp.baixar(principal))
        except Exception as erro:  # documento fora do modelo não impede os outros
            self.relatorio.avisos.append(f"{self._rotulo(grupo)}: {principal.nome} não foi lido ({erro}).")
            return DadosDocumento()

    @staticmethod
    def _rotulo(grupo: Grupo) -> str:
        return f"{grupo.numero:03d}/{grupo.ano % 100:02d} ({grupo.cliente or 'sem pasta de cliente'})"

    def importar(self, grupo: Grupo, linha: LinhaControle | None) -> None:
        proposta = Proposta.objects.filter(empresa=self.empresa, ano=grupo.ano, numero=grupo.numero).first()
        if proposta is not None and not proposta.importada:
            self.relatorio.conflitos.append(
                f"{self._rotulo(grupo)}: já existe {proposta.codigo} criada no sistema; o SharePoint não foi importado."
            )
            return
        principal = grupo.principal
        versao = f"{principal.id}:{principal.versao}"
        if proposta is None or proposta.sharepoint_versao != versao:
            dados = self._dados(grupo)
            nova = proposta is None
            proposta = self._gravar(proposta, grupo, linha, dados, versao)
            if nova:
                self.relatorio.criadas += 1
            else:
                self.relatorio.atualizadas += 1
                registrar(proposta, EventoProposta.EDITADA, f"Atualizada pelo SharePoint: {principal.nome}.")
        else:
            self.relatorio.sem_mudanca += 1
        if proposta.responsavel_id is None and proposta.assinada_por:
            # O usuário pode ter entrado no sistema depois da importação.
            proposta.responsavel = self._responsavel(proposta.assinada_por)
            if proposta.responsavel:
                proposta.save(update_fields=["responsavel", "atualizada_em"])
        self._arquivos(proposta, grupo)
        self._etapa(proposta, grupo)

    def _gravar(self, proposta, grupo: Grupo, linha, dados: DadosDocumento, versao: str) -> Proposta:
        nova = proposta is None
        linha = linha or LinhaControle("", "", None, "")
        cliente = self._cliente(grupo.cliente or linha.cliente or dados.cliente) if nova else proposta.cliente
        data = dados.data or linha.data or (grupo.principal.criado.date() if grupo.principal.criado else self.hoje)
        if nova:
            proposta = Proposta(
                empresa=self.empresa, ano=grupo.ano, numero=grupo.numero, cliente=cliente, etapa=ENVIADA,
                probabilidade=PROBABILIDADE_PADRAO[ENVIADA], enviada_em=data,
                responsavel=self._responsavel(dados.assinante or linha.enviado_por),
            )
        proposta.revisao = grupo.revisao
        proposta.titulo = (dados.titulo or linha.servico or grupo.principal.nome.rsplit(".", 1)[0])[:200]
        proposta.servico = linha.servico[:150] or proposta.servico
        proposta.validade = dados.validade or proposta.validade or data + timedelta(days=self.empresa.validade_dias)
        proposta.contato = self._contato(proposta.cliente, dados.contato) or proposta.contato
        proposta.condicoes_pagamento = dados.pagamento or proposta.condicoes_pagamento
        proposta.prazo_entrega = dados.prazo or proposta.prazo_entrega
        proposta.escopo = dados.escopo or proposta.escopo
        proposta.assinada_por = (dados.assinante or linha.enviado_por)[:150] or proposta.assinada_por
        notas = [f"Solicitante (planilha de controle): {linha.solicitante}." if linha.solicitante else ""]
        if dados.valor_a_conferir:
            notas.append("Mais de um valor no documento: confira o total.")
        proposta.observacoes = "\n".join(filter(None, notas)) or proposta.observacoes
        proposta.sharepoint_item, proposta.sharepoint_versao = grupo.principal.id, versao
        proposta.sharepoint_pasta_url = grupo.pasta_url or proposta.sharepoint_pasta_url
        proposta.importada_em = timezone.now()
        proposta.save()

        if dados.valor is None:
            self.relatorio.avisos.append(f"{self._rotulo(grupo)}: valor não encontrado no documento; ficou R$ 0,00.")
        elif dados.valor_a_conferir:
            self.relatorio.avisos.append(f"{self._rotulo(grupo)}: mais de um valor no documento; usado {dados.valor}.")
        valor = dados.valor if dados.valor is not None else ZERO
        itens = list(proposta.itens.all())
        if nova or (len(itens) == 1 and itens[0].item_catalogo_id is None):
            item = itens[0] if itens else ItemProposta(proposta=proposta)
            item.descricao, item.quantidade, item.preco_unitario = proposta.titulo[:255], Decimal("1"), valor
            item.save()
        elif dados.valor is not None and sum((i.total for i in itens), ZERO) != valor:
            self.relatorio.avisos.append(f"{self._rotulo(grupo)}: itens editados no sistema; o valor do documento ({valor}) não foi aplicado.")
        recalcular(proposta)
        if nova:
            registrar(proposta, EventoProposta.CRIADA, f"Importada do SharePoint: {grupo.principal.caminho}.", para=ENVIADA)
        return proposta

    def importar_so_planilha(self, ano: int, numero: int, linha: LinhaControle) -> None:
        """Linha da planilha de controle sem documento no SharePoint: a proposta entra assim mesmo, sem valor."""
        proposta = Proposta.objects.filter(empresa=self.empresa, ano=ano, numero=numero).first()
        rotulo = f"{numero:03d}/{ano % 100:02d} ({linha.cliente or 'sem cliente'})"
        if proposta is not None:
            if not proposta.importada:
                self.relatorio.conflitos.append(f"{rotulo}: já existe {proposta.codigo} criada no sistema; a planilha foi ignorada.")
            else:
                self.relatorio.sem_mudanca += 1
            return
        data = linha.data or date(ano, 1, 1)
        validade = data + timedelta(days=self.empresa.validade_dias)
        proposta = Proposta.objects.create(
            empresa=self.empresa, ano=ano, numero=numero, cliente=self._cliente(linha.cliente),
            titulo=(linha.servico or f"Proposta {numero:03d}/{ano % 100:02d}")[:200], servico=linha.servico[:150],
            etapa=ENVIADA, probabilidade=PROBABILIDADE_PADRAO[ENVIADA], enviada_em=data, validade=validade,
            observacoes="Só na planilha de controle: sem documento da proposta no SharePoint, valor desconhecido."
            + (f"\nSolicitante (planilha de controle): {linha.solicitante}." if linha.solicitante else ""),
            sharepoint_item=f"planilha:{ano}:{numero}", sharepoint_versao="planilha", importada_em=timezone.now(),
        )
        registrar(proposta, EventoProposta.CRIADA, "Importada da planilha de controle (sem documento no SharePoint).", para=ENVIADA)
        self.relatorio.criadas += 1
        self.relatorio.avisos.append(f"{rotulo}: só na planilha de controle; entrou sem valor.")
        if validade < self.hoje:
            proposta.etapa, proposta.probabilidade, proposta.fechada_em = EXPIRADA, 0, validade
            proposta.save(update_fields=["etapa", "probabilidade", "fechada_em", "atualizada_em"])
            self.relatorio.expiradas += 1

    def atualizar_acervo(self) -> None:
        """Grava o catálogo de arquivos vistos nesta rodada e apaga os que sumiram do SharePoint."""
        existentes = {a.item_id: a for a in ArquivoSharePoint.objects.filter(empresa=self.empresa)}
        propostas = {(p.ano, p.numero): p for p in Proposta.objects.filter(empresa=self.empresa).select_related("cliente")}
        por_cliente: dict[tuple[int, int], list[Proposta]] = defaultdict(list)
        for p in propostas.values():
            por_cliente[(p.ano, p.cliente_id)].append(p)
        novos, mudados, vistos = [], [], set()
        campos = ["nome", "caminho", "url", "pasta_raiz", "ano", "pasta_cliente", "cliente", "proposta", "categoria", "extensao", "tamanho", "modificado_em"]
        for item, raiz, ano, pasta_cliente in self.acervo:
            if item.id in vistos:
                continue
            vistos.add(item.id)
            proposta = None
            cliente = self._clientes.get(_normal(pasta_cliente)) if pasta_cliente else None
            m = ARQUIVO_PROPOSTA.match(item.nome)
            if m and ano and int(m["aa"]) == ano % 100:
                proposta = propostas.get((ano, int(m["numero"])))
            elif ano and cliente:
                # Anexo (pedido, contrato, NDA…) que cita o número de uma proposta do mesmo cliente e ano.
                proposta = next((p for p in por_cliente[(ano, cliente.pk)] if cita_numero(item.nome, p.numero, ano)), None)
            valores = {
                "nome": item.nome[:255], "caminho": item.caminho[:1000], "url": item.url[:1000], "pasta_raiz": raiz[:255],
                "ano": ano, "pasta_cliente": pasta_cliente[:255], "cliente": cliente or (proposta.cliente if proposta else None),
                "proposta": proposta, "categoria": classificar(item.nome), "extensao": item.nome.rsplit(".", 1)[-1].lower()[:10] if "." in item.nome else "",
                "tamanho": item.tamanho, "modificado_em": item.modificado,
            }
            atual = existentes.get(item.id)
            if atual is None:
                novos.append(ArquivoSharePoint(empresa=self.empresa, item_id=item.id, **valores))
            elif any(getattr(atual, c) != v for c, v in valores.items()):
                for c, v in valores.items():
                    setattr(atual, c, v)
                mudados.append(atual)
        ArquivoSharePoint.objects.bulk_create(novos, batch_size=500)
        ArquivoSharePoint.objects.bulk_update(mudados, campos, batch_size=500)
        ArquivoSharePoint.objects.filter(empresa=self.empresa).exclude(item_id__in=vistos).delete()
        self.relatorio.arquivos = len(vistos)

    def importar_carteira(self) -> None:
        """Refaz a carteira de contratos a partir da planilha de controle de contratos, se configurada."""
        from .carteira import ler_carteira

        caminho = self.empresa.sharepoint_controle_contratos.strip().strip("/")
        if not caminho:
            return
        pasta, _, nome = caminho.rpartition("/")
        try:
            planilha = next((i for i in self.sp.listar(pasta) if i.nome.lower() == nome.lower()), None)
        except NaoEncontrado:
            planilha = None
        if planilha is None:
            self.relatorio.avisos.append(f"Planilha de controle de contratos “{caminho}” não encontrada.")
            return
        try:
            linhas = ler_carteira(self.sp.baixar(planilha))
        except Exception as erro:  # planilha corrompida não impede o resto
            self.relatorio.avisos.append(f"{planilha.nome}: não foi possível ler ({erro}).")
            return
        ContratoCliente.objects.filter(empresa=self.empresa).delete()
        ContratoCliente.objects.bulk_create(
            ContratoCliente(empresa=self.empresa, cliente=self._cliente_do_contrato(l["cliente_nome"]), **l) for l in linhas
        )
        self.relatorio.contratos = len(linhas)

    def _cliente_do_contrato(self, nome: str) -> Cliente | None:
        """Cliente pelo nome exato (sem caixa e acento) ou, se for o único, pela primeira palavra ("Aché Laboratórios" → ACHÉ)."""
        chave = _normal(nome)
        if chave in self._clientes:
            return self._clientes[chave]
        primeira = chave.split(" ")[0] if chave else ""
        candidatos = {c.pk: c for k, c in self._clientes.items() if primeira and k.split(" ")[0] == primeira}
        return next(iter(candidatos.values())) if len(candidatos) == 1 else None

    def _arquivos(self, proposta: Proposta, grupo: Grupo) -> None:
        atuais = []
        for rev, item in grupo.documentos:
            atuais.append((item, ArquivoProposta.DOCUMENTO, rev))
        atuais += [(item, ArquivoProposta.PEDIDO, None) for item in grupo.pedidos]
        for item, tipo, rev in atuais:
            ArquivoProposta.objects.update_or_create(
                proposta=proposta, item_id=item.id,
                defaults={
                    "tipo": tipo, "nome": item.nome[:255], "caminho": item.caminho[:1000], "url": item.url[:1000],
                    "revisao": rev, "tamanho": item.tamanho, "modificado_em": item.modificado,
                },
            )
        proposta.arquivos.exclude(item_id__in=[item.id for item, _, _ in atuais]).delete()

    def _etapa(self, proposta: Proposta, grupo: Grupo) -> None:
        """Pedido de compra → ganha; sem pedido e validade vencida → expirada. Só sobre o que a importação definiu."""
        if proposta.etapa not in (ENVIADA, EXPIRADA):
            return
        origem = proposta.etapa
        if grupo.pedidos:
            primeiro = min(grupo.pedidos, key=lambda p: p.criado or timezone.now())
            proposta.etapa, proposta.probabilidade = GANHA, PROBABILIDADE_PADRAO[GANHA]
            proposta.fechada_em = max(timezone.localtime(primeiro.criado).date() if primeiro.criado else self.hoje, proposta.enviada_em or date.min)
            texto = f"Pedido de compra no SharePoint: {primeiro.nome}."
            self.relatorio.ganhas += 1
        elif origem == ENVIADA and proposta.validade and proposta.validade < self.hoje:
            proposta.etapa, proposta.probabilidade, proposta.fechada_em = EXPIRADA, 0, proposta.validade
            texto = f"Validade de {proposta.validade:%d/%m/%Y} vencida sem pedido de compra no SharePoint."
            self.relatorio.expiradas += 1
        else:
            return
        proposta.save(update_fields=["etapa", "probabilidade", "fechada_em", "atualizada_em"])
        registrar(proposta, EventoProposta.ETAPA, f"{ROTULOS_ETAPA[origem]} → {ROTULOS_ETAPA[proposta.etapa]}. {texto}", de=origem, para=proposta.etapa)


# Uma execução sem fim há mais que isto é considerada abandonada (processo derrubado no meio).
LIMITE_EM_ANDAMENTO = timedelta(minutes=120)  # a primeira carga do histórico leva dezenas de minutos


class SincronizacaoEmAndamento(SharePointErro):
    pass


def encerrar_interrompidas() -> int:
    """Rodadas sem fim quando o sistema sobe foram derrubadas junto com o processo: ficam como interrompidas.

    O que já foi gravado fica (cada proposta grava sozinha); a próxima rodada continua de onde parou.
    """
    return ExecucaoSincronizacao.objects.filter(fim__isnull=True).update(
        fim=timezone.now(), erro="Interrompida: o sistema foi reiniciado durante a sincronização."
    )


def em_andamento(empresa: Empresa) -> ExecucaoSincronizacao | None:
    return ExecucaoSincronizacao.objects.filter(
        empresa=empresa, fim__isnull=True, inicio__gte=timezone.now() - LIMITE_EM_ANDAMENTO
    ).first()


def sincronizar(
    empresa: Empresa,
    cliente: SharePointClient | None = None,
    simular: bool = False,
    hoje: date | None = None,
    usuario=None,
) -> Relatorio:
    """Importa ou atualiza as propostas da empresa a partir do SharePoint e registra a execução.

    ``simular`` desfaz tudo no fim. Recusa (``SincronizacaoEmAndamento``) se já houver uma rodando.
    """
    with transaction.atomic():
        Empresa.objects.select_for_update().filter(pk=empresa.pk).first()
        if (atual := em_andamento(empresa)) is not None:
            raise SincronizacaoEmAndamento(
                f"Já há uma sincronização em andamento desde {timezone.localtime(atual.inicio):%H:%M}; aguarde ela terminar."
            )
        execucao = ExecucaoSincronizacao.objects.create(empresa=empresa, usuario=usuario, simulacao=simular)
    try:
        relatorio = _sincronizar(empresa, cliente, simular, hoje)
    except Exception as erro:
        execucao.erro, execucao.fim = str(erro)[:2000], timezone.now()
        execucao.save(update_fields=["erro", "fim"])
        raise
    for campo in ("criadas", "atualizadas", "sem_mudanca", "ganhas", "expiradas", "arquivos", "conflitos", "avisos"):
        setattr(execucao, campo, getattr(relatorio, campo))
    execucao.fim = timezone.now()
    execucao.save()
    return relatorio


def _sincronizar(empresa: Empresa, cliente: SharePointClient | None, simular: bool, hoje: date | None) -> Relatorio:
    cliente = cliente or SharePointClient.criar(empresa.sharepoint_site)
    sinc = Sincronizacao(empresa, cliente, hoje)
    grupos, controle = sinc.listar()
    # Simulação: tudo numa transação desfeita no fim. Sincronização: cada proposta grava sozinha,
    # para a primeira carga (milhares de propostas) não segurar uma transação por dezenas de minutos.
    with transaction.atomic() if simular else nullcontext():
        for chave in sorted(grupos):
            try:
                with transaction.atomic():
                    sinc.importar(grupos[chave], controle.get(chave))
            except Exception as erro:  # uma proposta com problema não impede as outras
                sinc.relatorio.avisos.append(f"{Sincronizacao._rotulo(grupos[chave])}: não importada ({erro}).")
        for chave in sorted(set(controle) - set(grupos)):
            try:
                with transaction.atomic():
                    sinc.importar_so_planilha(*chave, controle[chave])
            except Exception as erro:
                sinc.relatorio.avisos.append(f"{chave[1]:03d}/{chave[0] % 100:02d}: linha da planilha não importada ({erro}).")
        with transaction.atomic():
            sinc.atualizar_acervo()
        with transaction.atomic():
            sinc.importar_carteira()
        if not simular:
            from .documentos import atualizar_documentos  # import local: documentos usa o cliente do SharePoint

            sinc.relatorio.avisos += atualizar_documentos(empresa, cliente)
        if simular:
            transaction.set_rollback(True)
        else:
            Empresa.objects.filter(pk=empresa.pk).update(sincronizado_em=timezone.now())
    return sinc.relatorio
