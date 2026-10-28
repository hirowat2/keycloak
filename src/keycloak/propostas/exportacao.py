"""Exportação de relatórios para Excel (.xlsx)."""

import io
from datetime import date
from decimal import Decimal

from django.http import HttpResponse

FORMATO_MOEDA = '#,##0.00;[Red]-#,##0.00'


def planilha_xlsx(nome_arquivo: str, abas: list[tuple[str, list[str], list[list[object]]]]) -> HttpResponse:
    """Gera um .xlsx com uma aba por relatório: (título, cabeçalho, linhas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    livro.remove(livro.active)
    for titulo, cabecalho, linhas in abas:
        aba = livro.create_sheet(titulo[:31])
        aba.append(cabecalho)
        for celula in aba[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="065F46")
        for linha in linhas:
            aba.append([float(v) if isinstance(v, Decimal) else v for v in linha])
            for celula, valor in zip(aba[aba.max_row], linha):
                if isinstance(valor, Decimal):
                    celula.number_format = FORMATO_MOEDA
                elif isinstance(valor, date):
                    celula.number_format = "DD/MM/YYYY"
        for indice, _ in enumerate(cabecalho, start=1):
            largura = max(
                (len(str(c.value)) for c in aba[get_column_letter(indice)] if c.value is not None),
                default=10,
            )
            aba.column_dimensions[get_column_letter(indice)].width = min(max(largura + 2, 12), 60)
        aba.freeze_panes = "A2"

    saida = io.BytesIO()
    livro.save(saida)
    resposta = HttpResponse(
        saida.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resposta["Content-Disposition"] = f'attachment; filename="{nome_arquivo}.xlsx"'
    return resposta
