"""Carteira de contratos a partir da planilha de controle de contratos do SharePoint.

A planilha (ex.: "Contratos de Clientes/Controle de Contratos.xlsx" no site da Pharma) tem uma
linha por cliente e sistema, com vencimento, aditivos, link e valores. A situação de cada linha é
a cor da célula do cliente, explicada pela legenda "Status:" no topo (Parado, Em andamento,
Encerrado, Descomissionamento). A planilha é a fonte: a cada sincronização a carteira é refeita.
"""

import io
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from .models import ContratoCliente

SITUACAO_DA_LEGENDA = {
    "parado": ContratoCliente.PARADO,
    "em andamento": ContratoCliente.ANDAMENTO,
    "encerrado": ContratoCliente.ENCERRADO,
    "descomissionamento": ContratoCliente.DESCOMISSIONADO,
    "descomissionado": ContratoCliente.DESCOMISSIONADO,
}
# Início do nome normalizado da coluna → campo.
COLUNAS = [
    ("cliente ativo", "ativo"),
    ("cliente", "cliente_nome"),
    ("sistema", "sistema"),
    ("responsavel", "responsavel"),
    ("contrato/nda/pedido de compra atual em nome", "em_nome_de"),
    ("contrato/nda assinado", "assinado"),
    ("venc", "vencimento"),
    ("data de fechamento", "fechamento"),
    ("observacao do contrato", "historico"),
    ("link", "link"),
    ("observacao", "observacao"),
    ("valor licenciamento", "valor_licenciamento"),
    ("valor hospedagem", "valor_hospedagem"),
    ("valot total", "valor_total"),
    ("valor total", "valor_total"),
]


def _normal(texto) -> str:
    base = unicodedata.normalize("NFKD", str(texto or "").replace("\xa0", " "))
    return " ".join(base.encode("ascii", "ignore").decode().lower().split())


def _cor(celula) -> str | None:
    preenchimento = celula.fill
    if not preenchimento or not preenchimento.fill_type:
        return None
    cor = preenchimento.fgColor
    if cor.type == "rgb":
        return str(cor.rgb)
    return f"{cor.type}:{cor.theme if cor.type == 'theme' else cor.indexed}:{round(cor.tint or 0, 2)}"


def _data(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(str(valor or "").strip(), formato).date()
        except ValueError:
            continue
    return None


def _sim_nao(valor) -> bool | None:
    texto = _normal(valor)
    return True if texto.startswith("sim") else False if texto.startswith("nao") else None


def _decimal(valor) -> Decimal | None:
    if valor in (None, ""):
        return None
    try:
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _aba(livro):
    """A aba com o cabeçalho Cliente/Sistema (de preferência a chamada "Controle")."""
    abas = sorted(livro.worksheets, key=lambda a: a.title.strip().lower() != "controle")
    for aba in abas:
        for linha in aba.iter_rows(min_row=1, max_row=15):
            nomes = [_normal(c.value) for c in linha]
            if "cliente" in nomes and "sistema" in nomes:
                return aba, linha[0].row
    return None, None


def ler_carteira(conteudo: bytes) -> list[dict]:
    """Linhas da carteira, com os campos de ``ContratoCliente``."""
    from openpyxl import load_workbook

    aba, cabecalho = _aba(load_workbook(io.BytesIO(conteudo), data_only=True))
    if aba is None:
        return []
    legenda = {}
    for linha in aba.iter_rows(min_row=1, max_row=cabecalho):
        if any(_normal(c.value).startswith("status") for c in linha):
            legenda = {_cor(c): SITUACAO_DA_LEGENDA[_normal(c.value)] for c in linha if _normal(c.value) in SITUACAO_DA_LEGENDA}
    colunas = {}
    for celula in aba[cabecalho]:
        nome = _normal(celula.value)
        campo = next((c for prefixo, c in COLUNAS if nome.startswith(prefixo)), None)
        if campo and campo not in colunas.values():
            colunas[celula.column - 1] = campo

    linhas = []
    for linha in aba.iter_rows(min_row=cabecalho + 1, max_row=aba.max_row):
        valores = {campo: linha[i].value if i < len(linha) else None for i, campo in colunas.items()}
        nome = str(valores.get("cliente_nome") or "").strip()
        if not nome:
            continue
        celula_link = next((linha[i] for i, campo in colunas.items() if campo == "link" and i < len(linha)), None)
        link = celula_link.hyperlink.target if celula_link is not None and celula_link.hyperlink else ""
        indice_cliente = next(i for i, campo in colunas.items() if campo == "cliente_nome")
        linhas.append(
            {
                "cliente_nome": nome[:200],
                "sistema": str(valores.get("sistema") or "").strip()[:255],
                "situacao": legenda.get(_cor(linha[indice_cliente]), ""),
                "ativo": _sim_nao(valores.get("ativo")),
                "responsavel": str(valores.get("responsavel") or "").strip(),
                "em_nome_de": str(valores.get("em_nome_de") or "").strip()[:150],
                "assinado": str(valores.get("assinado") or "").strip()[:100],
                "vencimento": _data(valores.get("vencimento")),
                "fechamento": _data(valores.get("fechamento")),
                "historico": str(valores.get("historico") or "").strip(),
                "observacao": str(valores.get("observacao") or "").strip(),
                "link": (link or (valores.get("link") if str(valores.get("link") or "").startswith("http") else ""))[:1000],
                "valor_licenciamento": _decimal(valores.get("valor_licenciamento")),
                "valor_hospedagem": _decimal(valores.get("valor_hospedagem")),
                "valor_total": _decimal(valores.get("valor_total")),
                "linha": linha[0].row,
            }
        )
    return linhas
